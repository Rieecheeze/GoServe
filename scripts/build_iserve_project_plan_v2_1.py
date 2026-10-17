from __future__ import annotations

import sys
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


NAVY = "1F2A44"
LIGHT_BLUE = "EAF2F8"
LIGHT_YELLOW = "FFF4CC"
LIGHT_GRAY = "F3F5F7"
WHITE = "FFFFFF"
GREEN = "16803A"
ORANGE = "B45309"
RED = "B42318"

TITLE = "IServe: A Native Android Application for Campus Operations, Services, and Consultations"
VERSION = "Version 2.1"
DATE = "October 13, 2026"


def set_cell_shading(cell, fill: str) -> None:
    properties = cell._tc.get_or_add_tcPr()
    shading = properties.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        properties.append(shading)
    shading.set(qn("w:fill"), fill)


def set_cell_margins(cell, top: int = 90, start: int = 100, bottom: int = 90, end: int = 100) -> None:
    properties = cell._tc.get_or_add_tcPr()
    margins = properties.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        properties.append(margins)
    for edge, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = margins.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row) -> None:
    properties = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    properties.append(header)


def keep_table_row_together(row) -> None:
    properties = row._tr.get_or_add_trPr()
    if properties.find(qn("w:cantSplit")) is None:
        properties.append(OxmlElement("w:cantSplit"))


def set_table_width(table, total_inches: float) -> None:
    properties = table._tbl.tblPr
    table_width = properties.find(qn("w:tblW"))
    if table_width is None:
        table_width = OxmlElement("w:tblW")
        properties.insert(0, table_width)
    table_width.set(qn("w:type"), "dxa")
    table_width.set(qn("w:w"), str(round(total_inches * 1440)))
    layout = properties.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        properties.append(layout)
    layout.set(qn("w:type"), "fixed")


def set_cell_width(cell, width_inches: float) -> None:
    cell.width = Inches(width_inches)
    properties = cell._tc.get_or_add_tcPr()
    width = properties.find(qn("w:tcW"))
    if width is None:
        width = OxmlElement("w:tcW")
        properties.append(width)
    width.set(qn("w:type"), "dxa")
    width.set(qn("w:w"), str(round(width_inches * 1440)))


def set_run_font(run, size: float | None = None, bold: bool | None = None, color: str | None = None) -> None:
    run.font.name = "Helvetica"
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Helvetica")
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Helvetica")
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)


def add_page_field(paragraph) -> None:
    prefix = paragraph.add_run("Page ")
    set_run_font(prefix, 7.5, color="5F6673")
    begin_run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    begin_run._r.append(begin)
    instruction_run = paragraph.add_run()
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    instruction_run._r.append(instruction)
    separate_run = paragraph.add_run()
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    separate_run._r.append(separate)
    result_run = paragraph.add_run("1")
    end_run = paragraph.add_run()
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    end_run._r.append(end)
    for run in (begin_run, instruction_run, separate_run, result_run, end_run):
        set_run_font(run, 7.5, color="5F6673")


def configure_document(document: Document) -> None:
    section = document.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.75)
    section.bottom_margin = Inches(0.75)
    section.left_margin = Inches(0.75)
    section.right_margin = Inches(0.75)

    normal = document.styles["Normal"]
    normal.font.name = "Helvetica"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Helvetica")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Helvetica")
    normal.font.size = Pt(9.5)
    normal.paragraph_format.space_after = Pt(4)
    normal.paragraph_format.line_spacing = 1.08

    heading = document.styles["Heading 1"]
    heading.font.name = "Helvetica"
    heading._element.rPr.rFonts.set(qn("w:ascii"), "Helvetica")
    heading._element.rPr.rFonts.set(qn("w:hAnsi"), "Helvetica")
    heading.font.size = Pt(15)
    heading.font.bold = True
    heading.font.color.rgb = RGBColor.from_string(NAVY)
    heading.paragraph_format.space_before = Pt(10)
    heading.paragraph_format.space_after = Pt(5)
    heading.paragraph_format.keep_with_next = True

    footer = section.footer
    table = footer.add_table(rows=1, cols=2, width=Inches(7.0))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.columns[0].width = Inches(5.9)
    table.columns[1].width = Inches(1.1)
    left = table.cell(0, 0).paragraphs[0]
    left.alignment = WD_ALIGN_PARAGRAPH.LEFT
    left_run = left.add_run(f"IServe Project Plan and Grounding Specification - v2.1 - {DATE}")
    set_run_font(left_run, 7.5, color="5F6673")
    right = table.cell(0, 1).paragraphs[0]
    right.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    add_page_field(right)
    for run in right.runs:
        set_run_font(run, 7.5, color="5F6673")


def add_label_paragraph(document: Document, label: str, text: str) -> None:
    paragraph = document.add_paragraph()
    label_color = {"[CONFIRMED]": GREEN, "[ASSUMPTION]": ORANGE, "[OPEN]": RED}[label]
    label_run = paragraph.add_run(label)
    set_run_font(label_run, 9.5, bold=True, color=label_color)
    text_run = paragraph.add_run(f" {text}")
    set_run_font(text_run, 9.5)


def add_bullet(document: Document, text: str) -> None:
    paragraph = document.add_paragraph(style="List Bullet")
    paragraph.paragraph_format.left_indent = Inches(0.2)
    paragraph.paragraph_format.first_line_indent = Inches(-0.12)
    run = paragraph.add_run(text)
    set_run_font(run, 9.5)


def add_metadata_table(document: Document) -> None:
    rows = (
        ("Project", TITLE),
        ("Document version", f"{VERSION}, written {DATE}"),
        ("Deployment context", "Sandbox demonstration for TIP Manila only"),
        ("Campus coverage", "Casal Campus and Arlegui Campus"),
        ("Front end", "Native Android application using Java and Android Studio (APK output)."),
        ("Authentication", "Firebase Authentication only"),
        ("Application backend", "Supabase PostgreSQL, Supabase Storage, Supabase Realtime, and Supabase Edge Functions"),
        ("Payments", "eGovPay sandbox when approved, with teller-confirmed cash as an alternative"),
        ("AI assistance", "Gemini for grounded guidance and temporary document analysis"),
        ("Deadline and team size", "Not provided [OPEN]"),
    )
    table = document.add_table(rows=1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_width(table, 7.0)
    table.columns[0].width = Inches(1.55)
    table.columns[1].width = Inches(5.45)
    headers = table.rows[0].cells
    headers[0].text = "Field"
    headers[1].text = "Value"
    set_repeat_table_header(table.rows[0])
    for cell in headers:
        set_cell_shading(cell, NAVY)
        set_cell_margins(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        for run in cell.paragraphs[0].runs:
            set_run_font(run, 9, bold=True, color=WHITE)
    set_cell_width(headers[0], 1.55)
    set_cell_width(headers[1], 5.45)
    for index, (field, value) in enumerate(rows):
        cells = table.add_row().cells
        cells[0].text = field
        cells[1].text = value
        set_cell_width(cells[0], 1.55)
        set_cell_width(cells[1], 5.45)
        if index % 2:
            for cell in cells:
                set_cell_shading(cell, LIGHT_GRAY)
        for cell in cells:
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
            for run in cell.paragraphs[0].runs:
                set_run_font(run, 8.8, bold=(cell is cells[0]))


def add_cover_and_legend(document: Document) -> None:
    title = document.add_paragraph()
    title.paragraph_format.space_after = Pt(7)
    title_run = title.add_run("IServe")
    set_run_font(title_run, 26, bold=True, color=NAVY)

    subtitle = document.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(4)
    subtitle_run = subtitle.add_run("A Native Android Application for Campus Operations, Services, and Consultations")
    set_run_font(subtitle_run, 15, bold=True, color=NAVY)

    descriptor = document.add_paragraph()
    descriptor.paragraph_format.space_after = Pt(10)
    descriptor_run = descriptor.add_run("Project Plan and Grounding Specification")
    set_run_font(descriptor_run, 11, color="5F6673")

    add_metadata_table(document)

    callout = document.add_table(rows=1, cols=1)
    callout.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = callout.cell(0, 0)
    set_cell_shading(cell, LIGHT_BLUE)
    set_cell_margins(cell, 120, 140, 120, 140)
    paragraph = cell.paragraphs[0]
    run = paragraph.add_run(
        "Purpose: a single source of truth for building and demonstrating IServe without inventing institutional facts."
    )
    set_run_font(run, 9.5, bold=True, color=NAVY)

    document.add_paragraph()
    intro = document.add_paragraph()
    intro_run = intro.add_run("Read this first. Every statement carries or sits under one of three labels.")
    set_run_font(intro_run, 9.5, bold=True)
    add_label_paragraph(document, "[CONFIRMED]", "Agreed project scope or technical contract. Treat it as binding for this plan.")
    add_label_paragraph(document, "[ASSUMPTION]", "A demo default chosen to keep implementation moving. Keep it visible and verify it before real deployment.")
    add_label_paragraph(document, "[OPEN]", "Information not yet issued or confirmed. Do not guess; use a visible placeholder until resolved.")


def add_section_0(document: Document) -> None:
    document.add_heading("0. Rules for AI models and readers using this document", level=1)

    rules = (
        "Use only the decisions and contracts stated in this plan. If a fact is not specified, identify the matching [OPEN] item in Section 15 instead of inventing it.",
        "Do not invent official TIP policies, fees, services, schedules, processing times, staff names, requirements, or approval rules. Seeded examples must be labeled as non-official demo data.",
        "Treat IServe as a sandbox demonstration for TIP Manila's Casal Campus and Arlegui Campus only. TIP Quezon City, outside institutions, and changes to official registrar grade records are out of scope.",
        "Use the exact role, office-status, request-status, and payment-status values defined in Sections 3, 5, and 6.",
        "Firebase is used only for authentication. Do not place application records or files in Firebase products.",
        "Use Supabase PostgreSQL for structured records, Supabase Storage for files, Supabase Realtime for in-app live updates, and Supabase Edge Functions for trusted server operations.",
        "Treat eGovPay callbacks and server-side reconciliation as authoritative. An Android app return, demo button, or client request cannot mark a payment as paid.",
        "Use teller-confirmed cash only through an authorized teller account. Keep payment status separate from request progress.",
        "Keep notifications in the application only. Do not add email notifications, device push, or Firebase Cloud Messaging.",
        "Gemini may answer only from approved knowledge records or the current temporary document. Do not retain chat history or temporary analysis files.",
        "Do not add new features merely because a technology supports them. Every professor-facing proposal feature listed in Section 5 is P1, and the agreed exclusions remain excluded.",
        "When citing or implementing this plan, preserve the [CONFIRMED], [ASSUMPTION], and [OPEN] distinction and quote the relevant section number.",
    )
    for rule in rules:
        add_bullet(document, rule)

    callout = document.add_table(rows=1, cols=1)
    callout.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = callout.cell(0, 0)
    set_cell_shading(cell, LIGHT_YELLOW)
    set_cell_margins(cell, 120, 140, 120, 140)
    paragraph = cell.paragraphs[0]
    lead = paragraph.add_run("Suggested preamble for another AI model: ")
    set_run_font(lead, 9.2, bold=True, color=NAVY)
    body = paragraph.add_run(
        "You are helping build IServe, a sandbox native Android application for TIP Manila's Casal and Arlegui campuses. "
        "Use this plan as the source of project facts. [CONFIRMED] is binding, [ASSUMPTION] is a visible demo default, "
        "and [OPEN] must not be guessed. Do not add features, institutions, official data, or integrations that are not stated here."
    )
    set_run_font(body, 9.2)


def add_subheading(document: Document, text: str) -> None:
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(6)
    paragraph.paragraph_format.space_after = Pt(3)
    paragraph.paragraph_format.keep_with_next = True
    run = paragraph.add_run(text)
    set_run_font(run, 11, bold=True, color=NAVY)


def add_body(document: Document, text: str) -> None:
    paragraph = document.add_paragraph()
    run = paragraph.add_run(text)
    set_run_font(run, 9.5)


def add_callout(document: Document, text: str, fill: str = LIGHT_BLUE) -> None:
    table = document.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.cell(0, 0)
    set_cell_shading(cell, fill)
    set_cell_margins(cell, 120, 140, 120, 140)
    run = cell.paragraphs[0].add_run(text)
    set_run_font(run, 9.2, bold=True, color=NAVY)


def add_standard_table(document: Document, headers: tuple[str, ...], rows: tuple[tuple[str, ...], ...], widths=None) -> None:
    table = document.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = widths is None
    if widths:
        set_table_width(table, sum(widths))
        for index, width in enumerate(widths):
            table.columns[index].width = Inches(width)
    header_cells = table.rows[0].cells
    set_repeat_table_header(table.rows[0])
    keep_table_row_together(table.rows[0])
    for index, text in enumerate(headers):
        header_cells[index].text = text
        if widths:
            set_cell_width(header_cells[index], widths[index])
        set_cell_shading(header_cells[index], NAVY)
        set_cell_margins(header_cells[index])
        for run in header_cells[index].paragraphs[0].runs:
            set_run_font(run, 8.4, bold=True, color=WHITE)
    for row_index, values in enumerate(rows):
        row = table.add_row()
        keep_table_row_together(row)
        cells = row.cells
        for column_index, value in enumerate(values):
            cells[column_index].text = value
            if widths:
                set_cell_width(cells[column_index], widths[column_index])
            cells[column_index].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
            set_cell_margins(cells[column_index])
            if row_index % 2:
                set_cell_shading(cells[column_index], LIGHT_GRAY)
            for run in cells[column_index].paragraphs[0].runs:
                set_run_font(run, 8.2, bold=(column_index == 0))


def add_section_1(document: Document) -> None:
    document.add_heading("1. Project overview", level=1)
    add_subheading(document, "1.1 Problem")
    add_body(
        document,
        "Students can lose time visiting several offices to ask about availability, submit requests, follow up on documents, or arrange consultations. Office personnel also need one place to publish availability, receive complete submissions, and respond without exposing records to unrelated users. [CONFIRMED]",
    )
    add_subheading(document, "1.2 Solution")
    add_body(
        document,
        "IServe is a native Android application for TIP Manila's Casal Campus and Arlegui Campus. Students can view office status and announcements, submit document requests and evidence, choose eGovPay or cash when a configured fee applies, track progress, receive verified PDFs, open inquiry threads, and book guidance or faculty consultations. Authorized staff and tellers use the same application with role-scoped access. [CONFIRMED]",
    )
    add_subheading(document, "1.3 Demonstration boundary")
    add_body(
        document,
        "The project demonstrates the proposed workflows with non-official data. It does not represent a production deployment, connect to an official student information system, or modify official registrar grade records. [CONFIRMED]",
    )
    add_subheading(document, "1.4 Out of scope")
    for item in (
        "TIP Quezon City, campuses other than Casal and Arlegui, and users from outside TIP Manila.",
        "Changes to official registrar grade records or direct integration with an official grade database.",
        "Public email and password registration, platform-admin screens, email notifications, and device push notifications.",
        "Automatic refunds or payment voids for cancelled or rejected paid requests.",
        "Persistent chatbot history, persistent uploaded-document history, typing indicators, and online presence.",
        "Web browser portals and cross-platform mobile frameworks.",
    ):
        add_bullet(document, item + " [CONFIRMED]")
    add_callout(
        document,
        "Demo content rule: office services, fees, schedules, requirements, and chatbot knowledge remain non-official demo data until verified by authorized TIP personnel.",
        LIGHT_YELLOW,
    )


def add_section_2(document: Document) -> None:
    document.add_heading("2. Confirmed P1 scope", level=1)
    add_body(
        document,
        "Every professor-facing proposal feature is required for the demonstration. P1 means the feature must work in the integrated demo; it does not mean the feature is approved for institutional production. [CONFIRMED]",
    )
    rows = (
        ("P1-01", "Office availability and announcements", "Staff publish a current office status, optional note, and announcements; students receive live in-app updates."),
        ("P1-02", "Document requests", "Students submit requests and private evidence, track status, and receive payment-gated verified PDF deliverables."),
        ("P1-03", "Payments", "Students choose eGovPay or cash. eGovPay uses an approved sandbox or labeled simulator; tellers confirm cash."),
        ("P1-04", "Guidance consultations", "Students book available guidance slots subject to configured capacity, cancellation, and rescheduling rules."),
        ("P1-05", "Faculty consultations", "Students book a CCS faculty member through schedules managed by the faculty member or authorized CCS secretary."),
        ("P1-06", "Inquiry threads", "Students and assigned staff exchange asynchronous messages grouped by conversation and timestamp."),
        ("P1-07", "Physical pickup QR", "A ready physical request receives a hashed, single-use pickup token with native Android camera scan and manual entry paths."),
        ("P1-08", "Gemini assistance", "The chatbot answers from approved knowledge or one current temporary file, with consent and no saved history."),
        ("P1-09", "Realtime notifications", "Supabase Realtime updates in-app status, announcements, messages, appointments, and notifications."),
        ("P1-10", "Native Android access", "The Java application installs from an APK and uses native layouts and widgets on Android phones."),
    )
    add_standard_table(document, ("ID", "Capability", "P1 demonstration contract"), rows, (0.65, 1.9, 4.45))


def add_section_3(document: Document) -> None:
    document.add_heading("3. Users, roles, and demo offices", level=1)
    add_subheading(document, "3.1 Authorization roles")
    role_rows = (
        ("student", "Verified TIP student or pre-created demo student", "Own profile, requests, payments, appointments, inquiry threads, files, and notifications; public office information."),
        ("staff", "Office personnel, secretary, faculty member, or guidance counselor", "Records assigned to the staff member's office or capability; permitted office content and schedules."),
        ("teller", "Cashiering teller", "Payment lookup and cash confirmation across configured demo offices; no authority to change unrelated request content."),
    )
    add_standard_table(document, ("Role", "Who", "Access boundary"), role_rows, (0.85, 2.25, 3.9))
    add_body(
        document,
        "Staff capability identifies whether a staff account acts as office personnel, secretary, faculty, or guidance counselor. The role remains staff so authorization policies have one stable role vocabulary. [CONFIRMED]",
    )
    add_subheading(document, "3.2 Demo offices")
    office_rows = (
        ("OFF-REG", "Registrar", "Casal Campus / Arlegui Campus", "Document requests, evidence review, verified deliverables, inquiries"),
        ("OFF-ACC", "Student Accounting", "Casal Campus / Arlegui Campus", "Accounting inquiries, fee-linked requests, announcements"),
        ("OFF-CASH", "Cashiering", "Casal Campus / Arlegui Campus", "Cash confirmation and payment inquiry demonstration"),
        ("OFF-GCC", "Guidance and Counseling Center", "Casal Campus / Arlegui Campus", "Guidance consultations and inquiry threads"),
        ("DEP-CCS", "College of Computer Studies secretary", "Casal Campus / Arlegui Campus", "Faculty consultation coordination, requests, and inquiries"),
    )
    add_standard_table(document, ("ID", "Office", "Demo coverage", "Demonstrated workflows"), office_rows, (0.8, 1.65, 1.65, 2.9))
    add_body(
        document,
        "The office list supports the agreed demonstration only. It does not claim that each service, schedule, or fee is officially offered at both campuses. [ASSUMPTION]",
    )
    add_subheading(document, "3.3 Office availability")
    status_rows = (
        ("Open", "Office is accepting the demonstrated service or inquiry."),
        ("On Lunch Break", "Office is temporarily unavailable during a manually indicated break."),
        ("Cut-Off Reached", "Office has reached the current service or queue cut-off."),
        ("Temporarily Closed", "Office is unavailable until staff update the status."),
    )
    add_standard_table(document, ("Status", "Meaning"), status_rows, (1.7, 5.3))
    add_body(
        document,
        "Authorized staff update office status manually. Each update may include a note and must record the actor and timestamp. [CONFIRMED]",
    )


def add_section_4(document: Document) -> None:
    document.add_heading("4. Payment methods and controls", level=1)
    add_subheading(document, "4.1 Shared payment contract")
    for item in (
        "Before checkout, the student chooses payment method egovpay or cash.",
        "The configured request-type fee is authoritative. The Android client cannot supply or modify the amount charged.",
        "Payment status is stored separately from request progress using pending, paid, failed, expired, or cancelled.",
        "A payment record and its paymentEvents remain auditable even when the linked request later changes status.",
        "A configured, clearly labeled test fee is required before any payment demonstration.",
    ):
        add_bullet(document, item + " [CONFIRMED]")
    add_subheading(document, "4.2 eGovPay flow")
    egovpay_rows = (
        ("1", "Student selects eGovPay for a fee-based request.", "The Android client sends only the request identifier over HTTPS to a trusted Edge Function."),
        ("2", "Edge Function loads the configured fee.", "Server creates the transaction and HMAC signature using secrets unavailable to the Android client."),
        ("3", "Student completes the sandbox flow.", "Official sandbox is used after access approval; otherwise the app uses a clearly labeled simulator."),
        ("4", "Provider sends a callback.", "Callback is authenticated, stored idempotently in paymentEvents, and cannot be applied twice."),
        ("5", "Server reconciles the transaction.", "Provider transaction status is checked before the payment becomes paid."),
        ("6", "Request continues.", "A confirmed paid record allows the request to enter or continue Processing."),
    )
    add_standard_table(document, ("Step", "User or provider event", "Trusted system action"), egovpay_rows, (0.55, 2.45, 4.0))
    add_callout(
        document,
        "Authoritative payment rule: the eGovPay callback and server-side reconciliation determine payment status. Returning to the Android app or tapping a button cannot mark a payment paid.",
        LIGHT_YELLOW,
    )
    add_subheading(document, "4.3 Cash flow")
    cash_rows = (
        ("1", "Student selects cash and receives a request reference with the configured amount due."),
        ("2", "Student pays through the demonstrated cashier process outside the Android application."),
        ("3", "An authenticated teller searches the reference and confirms the exact pending amount."),
        ("4", "The Edge Function records actor and time, rejects duplicate confirmation, marks payment paid, and advances eligible work."),
    )
    add_standard_table(document, ("Step", "Cash demonstration behavior"), cash_rows, (0.55, 6.45))
    add_subheading(document, "4.4 Cancellation, rejection, and provider limits")
    add_body(
        document,
        "Cancellation or rejection requires a reason. If the request has a paid record, the payment remains paid and the request enters manual office follow-up. The demonstration does not promise an automatic refund or void. [CONFIRMED]",
    )
    add_body(
        document,
        "The issued eGovPay base URL, access token, settlement template UUID, callback schema, and provider status values remain [OPEN] until sandbox approval. The provider-side already-paid control described for the presentation remains unverified and must not be presented as an official capability until tested. [OPEN]",
    )


def add_section_5(document: Document) -> None:
    document.add_heading("5. Functional requirements", level=1)
    add_body(document, "All requirements in this section are P1 for the sandbox demonstration. [CONFIRMED]")
    rows = (
        ("FR-01", "TIP Google sign-in", "Firebase Google sign-in accepts a verified @tip.edu.ph account. The bootstrap function automatically creates a student profile."),
        ("FR-02", "Demo sign-in", "Firebase email and password accepts only pre-created demo accounts. The app exposes no public registration path."),
        ("FR-03", "Role authorization", "student, staff, and teller routes and records are enforced through server checks and Firebase-token-based RLS."),
        ("FR-04", "Office directory", "Students browse the five demo offices and see campus-labeled demo services, availability, notes, and announcements."),
        ("FR-05", "Office status", "Authorized staff select one exact office status and may add a note. The system records actor and time."),
        ("FR-06", "Announcements", "Authorized staff publish office announcements. Relevant students receive in-app Realtime updates."),
        ("FR-07", "Document request", "A student selects a configured request type, enters notes, uploads valid private evidence, and receives a reference number."),
        ("FR-08", "Request tracker", "Students and authorized staff see request progress and history using the exact Section 6 statuses."),
        ("FR-09", "Payment choice", "The student selects egovpay or cash before checkout. The server loads the configured fee."),
        ("FR-10", "eGovPay transaction", "An Edge Function creates the transaction and signature, processes idempotent callbacks, and reconciles provider status."),
        ("FR-11", "Cash confirmation", "Only teller accounts can confirm a pending cash payment. A second confirmation is rejected."),
        ("FR-12", "Verified PDF delivery", "Authorized staff upload a private PDF. The owning student downloads it only when request and payment conditions allow."),
        ("FR-13", "Inquiry threads", "Students send asynchronous messages to an office or assigned recipient. The app stores conversation, sender, assignment, and time."),
        ("FR-14", "Guidance consultation", "Students book, cancel, or reschedule guidance slots without exceeding capacity."),
        ("FR-15", "Faculty consultation", "Students book a selected CCS faculty member. Authorized faculty or the CCS secretary manages the schedule."),
        ("FR-16", "Pickup token", "A physical request that becomes ready receives a hashed, single-use pickup token for CameraX-based QR scan or manual entry."),
        ("FR-17", "Gemini knowledge answer", "The Edge Function sends approved knowledge to Gemini and labels the response as guidance, not an official determination."),
        ("FR-18", "Temporary analysis", "A consenting user submits one PDF, JPG, or PNG up to 10 MB. The system deletes the private temporary file immediately after analysis."),
        ("FR-19", "In-app notifications", "Supabase Realtime updates the notification list and affected screens. The project sends no email or device push."),
        ("FR-20", "Retention", "A cleanup function deletes request evidence and deliverables 30 days after the request enters a terminal status."),
        ("FR-21", "Cancellation and rejection", "Authorized actors provide a reason. Paid records remain intact and require manual office follow-up."),
        ("FR-22", "Native Android access", "The APK runs on Android phones with native layouts and widgets; manual token entry remains available when camera permission or access is unavailable."),
    )
    add_standard_table(document, ("ID", "Requirement", "Rule"), rows, (0.7, 1.55, 4.75))


def add_section_6(document: Document) -> None:
    document.add_heading("6. State and lifecycle contracts", level=1)
    add_subheading(document, "6.1 Request statuses")
    request_rows = (
        ("Submitted", "The student created the request and the system assigned a reference number.", "System after validated submission"),
        ("Awaiting Payment", "A configured fee is due and no paid payment record exists.", "System after payment method selection"),
        ("Processing", "Payment is confirmed or no fee applies, and the office is working on the request.", "System or authorized office staff"),
        ("Ready for Pickup/Download", "The office prepared a physical result or verified digital deliverable.", "Authorized office staff"),
        ("Completed", "The student received the physical result or met the digital completion rule.", "Authorized office staff or controlled completion action"),
        ("Cancelled", "The request stopped with a required cancellation reason.", "Student where allowed or authorized staff"),
        ("Rejected", "The office declined the request with a required reason.", "Authorized office staff"),
    )
    add_standard_table(document, ("Status", "Meaning", "Set by"), request_rows, (1.65, 3.55, 1.8))
    add_body(
        document,
        "Completed, Cancelled, and Rejected are terminal request statuses for retention timing. Every transition stores the previous status, new status, actor, timestamp, and required reason when applicable. [CONFIRMED]",
    )
    add_subheading(document, "6.2 Payment statuses")
    payment_rows = (
        ("pending", "A payment path exists but authoritative confirmation has not arrived."),
        ("paid", "A teller confirmation or reconciled provider transaction proves payment for the stored amount."),
        ("failed", "The provider reports a failed attempt or trusted processing cannot complete it."),
        ("expired", "The provider or simulator validity period ended without payment."),
        ("cancelled", "The payment attempt was cancelled without changing any separate paid record."),
    )
    add_standard_table(document, ("Status", "Meaning"), payment_rows, (1.2, 5.8))
    add_callout(
        document,
        "Request progress and payment status are separate state machines. A request cannot infer paid status from its own progress, and a payment cannot silently rewrite a terminal request.",
        LIGHT_BLUE,
    )
    add_subheading(document, "6.3 Pickup token lifecycle")
    pickup_rows = (
        ("Issued", "The server creates an opaque token, stores only its hash, links it to one ready physical request, and sets its validity."),
        ("Presented", "Authorized staff scan the QR with CameraX and a QR decoder after Android camera permission is granted, or enter the token manually."),
        ("Consumed", "The server matches the hash, checks request and office authority, records the actor and time, and completes the pickup once."),
        ("Rejected", "The server rejects an unknown, expired, already consumed, or unauthorized token without completing the request."),
    )
    add_standard_table(document, ("Stage", "Contract"), pickup_rows, (1.15, 5.85))


def add_section_7(document: Document) -> None:
    document.add_heading("7. Consultations and inquiry threads", level=1)
    add_subheading(document, "7.1 Guidance and faculty consultations")
    consultation_rows = (
        ("Guidance", "Guidance and Counseling Center", "Guidance counselor", "Configured slot start, duration, and capacity"),
        ("Faculty", "College of Computer Studies secretary", "Selected CCS faculty member or authorized CCS secretary", "Faculty-specific availability and configured capacity"),
    )
    add_standard_table(document, ("Type", "Owning office", "Schedule manager", "Capacity basis"), consultation_rows, (1.0, 2.0, 2.15, 1.85))
    for item in (
        "The server rejects a booking when the active booking count reaches the configured slot capacity.",
        "The student may cancel or reschedule within the demonstration rules. Each action records actor and time and generates an in-app notification.",
        "A reschedule checks the target slot before releasing the original booking so capacity cannot be exceeded.",
        "Schedules and capacity values remain labeled demo data until authorized TIP personnel verify them.",
    ):
        add_bullet(document, item + " [CONFIRMED]")
    add_subheading(document, "7.2 Asynchronous inquiry threads")
    add_body(
        document,
        "An inquiryThread groups messages for one student, one destination office, and an optional assigned recipient. Each message stores the sender and timestamp. Staff assignment can route a thread to office personnel, the guidance counselor, the CCS secretary, or a CCS faculty member. [CONFIRMED]",
    )
    for item in (
        "Students read only their own conversations.",
        "Staff read only threads assigned to them or their authorized office capability.",
        "Messages are asynchronous. The demonstration excludes typing indicators and presence.",
        "Realtime subscriptions update authorized thread views and create in-app notifications without widening database access.",
    ):
        add_bullet(document, item)


def add_section_8(document: Document) -> None:
    document.add_heading("8. Technology and architecture", level=1)
    add_subheading(document, "8.1 Responsibility map")
    architecture_rows = (
        ("Android client", "Java and Android SDK", "Native phone layouts, widgets, navigation, and activity lifecycles for students, staff, faculty, guidance, and tellers. The client never holds provider or Gemini secrets."),
        ("Build and distribution", "Android Studio and Gradle", "Builds an installable APK for demonstration on Android phones. Firebase and Supabase continue to host their existing backend services."),
        ("Client networking", "Java HTTPS client and Supabase Realtime", "Runs requests off the UI thread; sends Firebase ID tokens to Supabase Data API, Storage, and Edge Functions over HTTPS; subscribes only to authorized live updates."),
        ("Identity provider", "Firebase Authentication", "Google sign-in for verified @tip.edu.ph users and email/password sign-in for pre-created demo accounts only."),
        ("Structured data", "Supabase PostgreSQL", "Stores all application records, relationships, histories, and audit fields."),
        ("Authorization", "Supabase RLS with Firebase tokens", "Uses Supabase's supported Firebase third-party authentication flow to enforce user, office, and capability boundaries."),
        ("File storage", "Supabase Storage", "Stores public-read office/service assets and private profile, evidence, deliverable, and temporary analysis files."),
        ("Live updates", "Supabase Realtime", "Publishes authorized database changes for in-app notifications, status, announcements, appointments, and messages."),
        ("Trusted operations", "Supabase Edge Functions", "Bootstraps users, authorizes operations, calculates charges, handles provider calls, serves Gemini, and runs cleanup."),
        ("Online payment", "eGovPay API", "Provides the approved sandbox transaction and status flow. A labeled simulator follows the same internal contract when credentials are unavailable."),
        ("AI provider", "Gemini", "Produces grounded guidance from approved records or the current temporary document through an Edge Function."),
    )
    add_standard_table(document, ("Layer", "Technology", "Responsibility"), architecture_rows, (1.25, 1.55, 4.2))
    add_subheading(document, "8.2 Authentication and authorization flow")
    auth_rows = (
        ("1", "The Firebase Authentication Android SDK authenticates a Google or pre-created demo account and issues an identity token."),
        ("2", "The Android client obtains a current Firebase ID token and presents it to Supabase APIs over HTTPS using the supported third-party authentication integration."),
        ("3", "A bootstrap Edge Function verifies identity, provisions a verified TIP Google user as student, or loads the pre-created demo profile."),
        ("4", "PostgreSQL and Storage RLS evaluate the Firebase identity and application role for every protected operation."),
        ("5", "Realtime sends only changes the same authenticated user may read under the configured policies."),
    )
    add_standard_table(document, ("Step", "Flow"), auth_rows, (0.55, 6.45))
    add_subheading(document, "8.3 Edge Function contracts")
    function_rows = (
        ("bootstrap-user", "Verify Firebase identity, enforce account rules, provision verified TIP students, and return application role/capability."),
        ("create-egovpay-transaction", "Load the authoritative fee, create a transaction, generate the HMAC signature, and store the pending attempt."),
        ("egovpay-callback", "Authenticate and store each callback idempotently, then trigger or record reconciliation."),
        ("reconcile-egovpay", "Query the provider transaction endpoint and map the approved provider status to the internal payment status."),
        ("confirm-cash-payment", "Require teller role, verify a pending cash amount, reject duplicates, and record actor and time."),
        ("ask-gemini", "Apply consent, grounding, file, rate, and provider-failure rules without exposing the API key."),
        ("cleanup-files", "Delete temporary analysis files immediately and eligible request files 30 days after terminal status."),
    )
    add_standard_table(document, ("Function", "Trusted responsibility"), function_rows, (2.2, 4.8))
    add_body(
        document,
        "Technical references: Supabase Firebase third-party authentication, https://supabase.com/docs/guides/auth/third-party/firebase-auth; eGovPay API catalog, https://platforms.e.gov.ph/api-catalogs/egovpay; eGovPay FAQ, https://platforms.e.gov.ph/faqs. [CONFIRMED references]",
    )
    add_body(
        document,
        "Android references: Firebase Authentication for Android, https://firebase.google.com/docs/auth/android/start; Android network operations, https://developer.android.com/develop/connectivity/network-ops/connecting; CameraX architecture, https://developer.android.com/media/camera/camerax/architecture.",
    )


def add_section_9(document: Document) -> None:
    document.add_heading("9. Data model, files, and privacy", level=1)
    add_subheading(document, "9.1 Structured records")
    model_rows = (
        ("users", "id, firebaseUid, email, displayName, role, staffCapability, officeId, isDemo, consentAcceptedAt"),
        ("offices", "id, name, campusLabel, status, statusNote, statusBy, statusAt, active"),
        ("announcements", "id, officeId, title, body, publishedBy, publishedAt, expiresAt"),
        ("requestTypes", "id, officeId, name, fee, deliveryMode, requirementsText, active, isDemo"),
        ("requests", "id, referenceNo, studentId, officeId, requestTypeId, status, notes, terminalAt, createdAt, updatedAt"),
        ("requestStatusEvents", "id, requestId, fromStatus, toStatus, reason, actorId, createdAt"),
        ("payments", "id, requestId, method, amount, status, providerReference, confirmedBy, confirmedAt, createdAt"),
        ("paymentEvents", "id, paymentId, providerEventId, eventType, payloadHash, status, receivedAt, reconciledAt"),
        ("files", "id, ownerId, requestId, bucket, objectPath, kind, mimeType, sizeBytes, expiresAt, createdAt"),
        ("appointments", "id, studentId, officeId, staffId, slotStart, slotEnd, status, createdAt, updatedAt"),
        ("slotSettings", "id, officeId, staffId, slotMinutes, capacity, availabilityData, active"),
        ("inquiryThreads", "id, studentId, officeId, assignedTo, subject, status, createdAt, updatedAt"),
        ("messages", "id, threadId, senderId, body, createdAt"),
        ("pickupTokens", "id, requestId, tokenHash, expiresAt, consumedAt, consumedBy"),
        ("notifications", "id, recipientId, type, title, body, relatedId, readAt, createdAt"),
        ("knowledgeArticles", "id, officeId, title, content, approved, approvedBy, updatedAt"),
        ("auditLogs", "id, actorId, action, targetType, targetId, metadata, createdAt"),
        ("appConfig", "demoMode, allowedEmailDomain, paymentSimulatorEnabled, providerConfigurationState"),
    )
    add_standard_table(document, ("Table", "Key fields"), model_rows, (1.65, 5.35))
    add_body(
        document,
        "Field names are implementation contracts for the plan and may be normalized during schema design without changing the stated data ownership, status values, or security rules. [ASSUMPTION]",
    )
    add_subheading(document, "9.2 Storage classes")
    storage_rows = (
        ("office-assets", "Public read", "Office and service images approved for the demo", "Staff write; public read"),
        ("profile-images", "Private", "User profile images", "Owner read/write with authorized staff read only where required"),
        ("request-files", "Private", "Student evidence and staff PDF deliverables", "Owning student and authorized office staff under RLS"),
        ("temporary-ai", "Private", "One current PDF, JPG, or PNG for analysis", "Requesting user and Edge Function only"),
    )
    add_standard_table(document, ("Bucket", "Visibility", "Content", "Access"), storage_rows, (1.25, 1.0, 2.15, 2.6))
    add_subheading(document, "9.3 File validation and retention")
    for item in (
        "Request evidence follows configured type and size rules before upload metadata is accepted.",
        "Verified deliverables use PDF and remain private. A student must own the request and meet payment and request-status gates before download.",
        "Temporary analysis accepts one PDF, JPG, or PNG up to 10 MB after the user accepts the non-sensitive-data consent notice.",
        "The Gemini flow deletes the temporary object immediately after analysis, including handled provider failures.",
        "The cleanup function deletes request evidence and deliverables 30 days after Completed, Cancelled, or Rejected, and records the cleanup result.",
    ):
        add_bullet(document, item + " [CONFIRMED]")
    add_subheading(document, "9.4 Privacy controls")
    add_body(
        document,
        "The demonstration collects only the records needed for the agreed workflows. RLS blocks cross-student and cross-office access. Secrets stay in server-side configuration. Audit records identify sensitive actions without storing provider secrets, document contents, or Gemini prompts and responses. [CONFIRMED]",
    )
    add_body(
        document,
        "The team must review the demonstration against the Data Privacy Act of 2012 and TIP requirements before any real deployment. This plan is a technical project document, not legal advice. [ASSUMPTION]",
    )


def add_section_10(document: Document) -> None:
    document.add_heading("10. Screen inventory", level=1)
    screen_rows = (
        ("Shared", "Sign in", "Google sign-in, pre-created demo email/password sign-in, blocked public registration, error states"),
        ("Student", "Consent and profile", "Privacy consent, non-sensitive AI notice, private profile image"),
        ("Student", "Office directory", "Campus labels, office status, note, announcements, demo services"),
        ("Student", "Office detail", "Request, inquiry, consultation, and current office information"),
        ("Student", "New request", "Request type, notes, evidence upload, delivery mode, configured fee, payment choice"),
        ("Student", "Payment", "eGovPay handoff and authoritative status refresh on app return, or cash instructions; separate payment status and retry where allowed"),
        ("Student", "My requests", "Reference, request status history, payment status, verified download, pickup QR"),
        ("Student", "Consultations", "Guidance and faculty availability, booking, cancellation, rescheduling"),
        ("Student", "Inquiry threads", "Conversation list, assigned destination, asynchronous messages"),
        ("Student", "Notifications", "Realtime in-app events with read state"),
        ("Staff", "Office dashboard", "Office status, note, announcements, queue totals, appointments, inquiries"),
        ("Staff", "Request queue", "Office-scoped filters and request details"),
        ("Staff", "Request detail", "Evidence review, status action, required reason, verified PDF upload, pickup action"),
        ("Staff", "Pickup scanner", "CameraX QR scan after camera permission, manual token fallback, used-token rejection"),
        ("Guidance", "Guidance schedule", "Slot configuration, capacity, bookings, cancellation and reschedule events"),
        ("Faculty", "Faculty schedule", "Own availability, bookings, inquiry assignment"),
        ("Secretary", "CCS coordination", "Faculty schedules and assigned CCS workflows within authorization"),
        ("Teller", "Cash payments", "Reference search, authoritative amount, pending state, confirmation, duplicate rejection"),
        ("Demo", "Payment simulator", "Clearly labeled provider simulator with the same internal statuses and callbacks"),
        ("Shared", "Gemini assistant", "Approved knowledge questions, temporary file consent and upload, guidance label, refusal and failure states"),
    )
    add_standard_table(document, ("Audience", "Screen", "Contents"), screen_rows, (1.0, 1.65, 4.35))


def add_section_11(document: Document) -> None:
    document.add_heading("11. Demonstration configuration", level=1)
    add_subheading(document, "11.1 Pre-created accounts")
    account_rows = (
        ("1", "student", "Student A", "student.a@iserve.demo", "Own records only"),
        ("2", "student", "Student B", "student.b@iserve.demo", "Own records only; used for cross-student denial"),
        ("3", "staff", "Registrar staff", "registrar.staff@iserve.demo", "Registrar"),
        ("4", "staff", "Student Accounting staff", "accounting.staff@iserve.demo", "Student Accounting"),
        ("5", "teller", "Cashiering teller", "cashier.teller@iserve.demo", "Cashiering payment confirmation"),
        ("6", "staff", "Guidance counselor", "guidance.staff@iserve.demo", "Guidance capability"),
        ("7", "staff", "CCS secretary", "ccs.secretary@iserve.demo", "Secretary capability"),
        ("8", "staff", "CCS faculty", "ccs.faculty@iserve.demo", "Faculty capability"),
    )
    add_standard_table(document, ("#", "Role", "Demo identity", "Suggested email", "Assignment"), account_rows, (0.35, 0.75, 1.5, 2.2, 2.2))
    add_body(
        document,
        "These addresses are suggested non-production demo identifiers. Passwords, provider secrets, and real student information do not appear in this plan. [ASSUMPTION]",
    )
    add_subheading(document, "11.2 Seed-data rules")
    for item in (
        "Seed only the five demo offices listed in Section 3.",
        "Label every sample service, requirement, fee, processing estimate, office schedule, and knowledge article as non-official demo data.",
        "Configure at least one free request and one fee-based request so both paths can be shown.",
        "Set a test fee in server-controlled requestTypes data before demonstrating eGovPay or cash.",
        "Keep paymentSimulatorEnabled visible in the interface whenever official sandbox credentials are unavailable.",
        "Create appointment availability for the guidance counselor and CCS faculty member without presenting those times as official schedules.",
    ):
        add_bullet(document, item + " [CONFIRMED]")


def add_section_12(document: Document) -> None:
    document.add_heading("12. Acceptance test plan", level=1)
    tests = (
        ("AT-01", "Verified TIP Google user signs in for the first time.", "Firebase identity is accepted and the bootstrap function provisions role student."),
        ("AT-02", "Pre-created demo account signs in with email/password.", "The matching demo profile and authorization load."),
        ("AT-03", "Unknown user attempts email/password registration.", "No public signup path exists and account creation is blocked."),
        ("AT-04", "Personal Google account signs in.", "Provisioning is denied and protected data remains unreadable."),
        ("AT-05", "Each student, office staff, faculty, guidance, secretary, and teller account opens its allowed screens.", "Each account sees only its authorized routes and actions."),
        ("AT-06", "Student A requests Student B's record or private file.", "Database and Storage RLS deny access."),
        ("AT-07", "Registrar staff requests Student Accounting records.", "Cross-office RLS denies access."),
        ("AT-08", "Staff changes office status and publishes an announcement.", "Actor/time are stored and authorized students receive Realtime updates."),
        ("AT-09", "Student opens an inquiry and staff replies.", "Messages remain in one authorized thread with sender, assignment, and timestamp."),
        ("AT-10", "User attempts presence or typing updates.", "No such feature or persistent field exists."),
        ("AT-11", "Student uploads valid private evidence.", "Validated file metadata is stored and only authorized users can download it."),
        ("AT-12", "Student uploads a disallowed type or oversized request file.", "Upload is rejected with a clear validation message."),
        ("AT-13", "Staff uploads a verified PDF for a paid request that is ready.", "The owning student receives a private signed download path."),
        ("AT-14", "Student requests a paid deliverable before payment confirmation.", "Download is denied."),
        ("AT-15", "A request remains terminal for 30 days.", "Cleanup deletes eligible evidence and deliverables and records the result."),
        ("AT-16", "Student starts eGovPay checkout after a test fee is configured.", "The Edge Function charges the exact server-loaded amount."),
        ("AT-17", "Android client submits a different amount.", "The server ignores or rejects it and uses the configured fee."),
        ("AT-18", "The same eGovPay callback arrives twice.", "paymentEvents records idempotency and payment is applied once."),
        ("AT-19", "Android app returns from eGovPay without an authoritative paid result.", "Payment remains pending until callback and reconciliation confirm it."),
        ("AT-20", "Provider reports failed, expired, or cancelled.", "Internal payment status maps correctly without false request completion."),
        ("AT-21", "Labeled simulator runs each supported payment outcome.", "Internal events and visible results match the sandbox contract."),
        ("AT-22", "Teller confirms a pending cash payment.", "Payment becomes paid with actor/time and eligible request work continues."),
        ("AT-23", "Student or staff account attempts cash confirmation.", "Authorization denies the operation."),
        ("AT-24", "Teller confirms an already paid cash record.", "Duplicate confirmation is rejected without a second event."),
        ("AT-25", "Authorized user cancels or rejects a request without a reason.", "Transition is rejected."),
        ("AT-26", "A paid request is cancelled or rejected with a reason.", "Paid record remains and the case is flagged for manual office follow-up."),
        ("AT-27", "Guidance or faculty slot reaches capacity.", "Another booking is rejected."),
        ("AT-28", "Student cancels and reschedules a consultation.", "Capacity remains correct and authorized parties receive in-app updates."),
        ("AT-29", "Staff scans a valid unused pickup QR.", "Token is consumed once and the request completes."),
        ("AT-30", "Staff scans or enters the same pickup token again.", "The reused token is rejected."),
        ("AT-31", "Student asks a question covered by an approved knowledgeArticle.", "Gemini answers from that content and labels the answer as guidance."),
        ("AT-32", "Student asks an unsupported institutional question.", "The assistant refuses to invent an answer and points to the relevant office."),
        ("AT-33", "User submits one valid temporary PDF, JPG, or PNG at or below 10 MB after consent.", "Analysis runs without storing chat history and the file is deleted immediately."),
        ("AT-34", "User skips consent, submits multiple files, wrong type, or more than 10 MB.", "Analysis is blocked before the provider call."),
        ("AT-35", "User exceeds the configured Gemini request limit.", "The Edge Function rate-limits the request."),
        ("AT-36", "Gemini provider fails.", "The app shows a safe error, saves no response history, and deletes any temporary file."),
        ("AT-37", "A tester installs the generated APK on an Android phone and signs in.", "The native app launches, shows the account's authorized screens, and accesses the existing backend."),
        ("AT-38", "Staff denies Android camera permission and opens pickup completion.", "Manual token entry remains available and uses the same server validation."),
    )
    add_standard_table(document, ("ID", "Scenario", "Expected result"), tests, (0.7, 3.05, 3.25))


def add_section_13(document: Document) -> None:
    document.add_heading("13. Dependency-ordered work plan", level=1)
    add_body(
        document,
        "The deadline and team allocation are not confirmed, so the schedule uses dependencies instead of dates. Every package is P1 because each proposal feature must work in the demonstration. [CONFIRMED]",
    )
    rows = (
        ("WP-1", "Foundation", "Android Studio Java project, Android SDK setup, APK build, Supabase project, Firebase project, configuration boundaries", "None"),
        ("WP-2", "Identity and RLS", "Firebase sign-in, bootstrap function, roles/capabilities, demo accounts, PostgreSQL and Storage policies", "WP-1"),
        ("WP-3", "Offices and content", "Demo offices, office status, announcements, request types, demo-data labels, Realtime subscriptions", "WP-2"),
        ("WP-4", "Requests and files", "Request creation, evidence validation, status history, staff queue, verified PDF delivery, retention metadata", "WP-2, WP-3"),
        ("WP-5", "Payments", "Server fee lookup, eGovPay transaction/callback/reconciliation, simulator, teller cash confirmation, paymentEvents", "WP-2, WP-4"),
        ("WP-6", "Consultations", "Guidance and faculty availability, capacity, booking, cancellation, rescheduling", "WP-2, WP-3"),
        ("WP-7", "Inquiries and notifications", "Inquiry threads, assignments, messages, in-app Realtime notification list", "WP-2, WP-3"),
        ("WP-8", "Pickup completion", "Hashed single-use token, QR presentation, CameraX scan and permission handling, manual entry, reuse rejection", "WP-4"),
        ("WP-9", "Gemini assistant", "Approved knowledge, consent, temporary file analysis, refusal, limits, provider failure, immediate deletion", "WP-2, WP-3, WP-4"),
        ("WP-10", "Integrated verification", "Acceptance tests AT-01 through AT-38, privacy checks, Android phone APK and activity lifecycle QA, demo script, failure-state rehearsal", "WP-1 through WP-9"),
    )
    add_standard_table(document, ("WP", "Package", "Includes", "Depends on"), rows, (0.65, 1.35, 4.1, 0.9))
    add_callout(
        document,
        "Scope warning: all proposal features remain P1. If delivery capacity becomes insufficient, the team must renegotiate scope with the professor instead of silently removing a feature.",
        LIGHT_YELLOW,
    )


def add_section_14(document: Document) -> None:
    document.add_heading("14. Risks and controls", level=1)
    rows = (
        ("R-01", "All proposal features are P1", "Integration may exceed available time.", "Build in dependency order, demonstrate thin complete flows early, and raise scope risk with the professor."),
        ("R-02", "eGovPay approval or credentials arrive late", "Official sandbox flow may be unavailable.", "Keep the provider adapter isolated and use a clearly labeled simulator with the same internal states."),
        ("R-03", "Provider callback details differ from assumptions", "Payment mapping or verification may fail.", "Keep schema and status mapping [OPEN], validate signatures, store idempotent events, and reconcile server-side."),
        ("R-04", "Android app return is mistaken for payment proof", "An unpaid request could appear paid.", "Allow only authenticated callback and reconciliation logic to set paid."),
        ("R-05", "Fee or service data looks official", "The demonstration could misrepresent TIP services.", "Label seed records as non-official demo data and require authorized verification before production use."),
        ("R-06", "Firebase and Supabase identity mapping is incorrect", "RLS could deny valid users or expose records.", "Test bootstrap claims, every role, cross-student denial, and cross-office denial before feature demos."),
        ("R-07", "Private files use public paths", "Evidence or deliverables could be exposed.", "Use private buckets, RLS, short-lived access, and negative download tests."),
        ("R-08", "Retention job does not run", "Files remain past the stated period.", "Log cleanup runs, test eligible dates, and expose failures to the project team."),
        ("R-09", "Payment and request states become coupled", "Cancellation or provider failure could corrupt progress history.", "Store separate state machines and connect them through audited server actions."),
        ("R-10", "Pickup token is copied or reused", "The wrong person could complete a pickup.", "Store only a hash, bind the token to one request and office, set validity, and reject reuse."),
        ("R-11", "Consultation capacity changes concurrently", "A slot could be overbooked.", "Use a transaction or server-side atomic capacity check."),
        ("R-12", "Gemini invents institutional answers", "Students may treat guidance as official.", "Ground only on approved records or the current file, refuse unsupported questions, and label every answer as guidance."),
        ("R-13", "Sensitive content is sent for analysis", "The free demonstration could process unsuitable data.", "Show explicit non-sensitive-data consent, validate the file, keep no chat history, and delete the file immediately."),
        ("R-14", "Realtime subscription exposes another office", "Live updates could bypass intended isolation.", "Apply the same Firebase-token-based RLS to Realtime reads and test unauthorized subscriptions."),
        ("R-15", "Android camera permission is denied or unavailable", "Staff cannot scan the pickup QR.", "Provide manual token entry with the same server validation."),
    )
    add_standard_table(document, ("ID", "Risk", "Impact", "Control"), rows, (0.55, 1.7, 1.75, 3.0))


def add_section_15(document: Document) -> None:
    document.add_heading("15. Open questions register", level=1)
    add_body(document, "These items remain unresolved. Do not answer them on the project owner's behalf. [OPEN]")
    rows = (
        ("Q-01", "What is the final deadline, and how many team members will implement the project?", "Schedule and assignment"),
        ("Q-02", "Which official demo services, fees, evidence requirements, processing notes, delivery modes, and office schedules may the team use?", "Seed data and presentation accuracy"),
        ("Q-03", "What eGovPay sandbox base URL, credential format, settlement template UUID, callback schema, signature details, and provider status values will be issued?", "Payment adapter"),
        ("Q-04", "Does the approved eGovPay sandbox expose the described provider-side already-paid demonstration control, and what result does it send?", "Demo script; capability must not be promised before testing"),
        ("Q-05", "Which Gemini model and project quota will be available for the non-sensitive demonstration?", "Rate limits and provider configuration"),
        ("Q-06", "Which office personnel may approve knowledgeArticles and verify the final demo content?", "Chatbot grounding and content governance"),
        ("Q-07", "What submission artifacts does the subject require in addition to the APK and this plan?", "Final project package"),
    )
    add_standard_table(document, ("ID", "Question", "Needed for"), rows, (0.65, 4.85, 1.5))


def add_section_16(document: Document) -> None:
    document.add_heading("16. Glossary", level=1)
    rows = (
        ("IServe", TITLE + ", the sandbox native Android application described in this plan."),
        ("Firebase Authentication", "The only Firebase service used by IServe; it handles Google and pre-created demo email/password identities."),
        ("Firebase third-party authentication", "Supabase's supported method for accepting Firebase identity tokens for Database, Storage, and Realtime access."),
        ("RLS", "Row Level Security policies that enforce which authenticated user may read or change each record."),
        ("Supabase PostgreSQL", "The database for all structured application records."),
        ("Supabase Storage", "Object storage for public office assets and private user, request, deliverable, and temporary AI files."),
        ("Supabase Realtime", "Live delivery of authorized database changes to the in-app interface."),
        ("Supabase Edge Functions", "Trusted server code for identity bootstrap, payment operations, Gemini requests, and cleanup."),
        ("eGovPay", "The payment API targeted for the approved sandbox demonstration."),
        ("Payment simulator", "A clearly labeled fallback that reproduces the app's internal payment outcomes when official sandbox access is unavailable."),
        ("HMAC signature", "A server-generated message authentication code used by the payment integration; its secret never reaches the Android client."),
        ("Idempotent callback", "A provider event that the server can receive more than once without applying the payment more than once."),
        ("Reconciliation", "A server-side status check against the provider transaction endpoint before or after applying an event."),
        ("Teller", "The role allowed to confirm a pending cash payment."),
        ("Request progress", "The separate request state: Submitted, Awaiting Payment, Processing, Ready for Pickup/Download, Completed, Cancelled, or Rejected."),
        ("Payment status", "The separate payment state: pending, paid, failed, expired, or cancelled."),
        ("Terminal request status", "Completed, Cancelled, or Rejected; the 30-day request-file retention period starts when the request enters one of these statuses."),
        ("Pickup token", "An opaque, hashed, single-use value that staff validate for one physical pickup."),
        ("knowledgeArticles", "Approved records that supply institutional demo knowledge to the Gemini assistant."),
        ("Temporary analysis", "One consenting user's PDF, JPG, or PNG, up to 10 MB, processed for the current answer and deleted immediately."),
        ("Grounded answer", "A Gemini response limited to approved knowledgeArticles or the current temporary document and labeled as guidance."),
        ("Non-official demo data", "Sample services, fees, schedules, requirements, and knowledge that TIP personnel have not verified for production."),
    )
    add_standard_table(document, ("Term", "Meaning in this project"), rows, (2.0, 5.0))


def build(output_path: Path) -> None:
    output_path.parent.mkdir(parents=True, exist_ok=True)
    document = Document()
    configure_document(document)
    add_cover_and_legend(document)
    add_section_0(document)
    add_section_1(document)
    add_section_2(document)
    add_section_3(document)
    add_section_4(document)
    add_section_5(document)
    add_section_6(document)
    add_section_7(document)
    add_section_8(document)
    add_section_9(document)
    add_section_10(document)
    add_section_11(document)
    add_section_12(document)
    add_section_13(document)
    add_section_14(document)
    add_section_15(document)
    add_section_16(document)
    document.save(output_path)


def main() -> None:
    if len(sys.argv) != 2:
        raise SystemExit("usage: build_iserve_project_plan.py <output.docx>")
    build(Path(sys.argv[1]))


if __name__ == "__main__":
    main()
